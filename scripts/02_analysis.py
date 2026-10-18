from __future__ import annotations

import json
import math
import re
import unicodedata
from pathlib import Path

import geopandas as gpd
import numpy as np
import pandas as pd
from openpyxl import load_workbook

EOV_EPSG = 23700
OUTPUT_CRS = "EPSG:4326"


def normalize_text(value: object) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.lower().strip()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def normalize_name(value: object) -> str:
    return normalize_text(value)


def normalize_ksh_code(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, (int, np.integer)):
        digits = str(int(value))
        return digits.zfill(5) if len(digits) < 5 else digits
    if isinstance(value, (float, np.floating)):
        if math.isnan(float(value)):
            return None
        if float(value).is_integer():
            digits = str(int(value))
            return digits.zfill(5) if len(digits) < 5 else digits
    digits = re.sub(r"\D", "", str(value))
    if not digits:
        return None
    if len(digits) < 5:
        digits = digits.zfill(5)
    return digits


def find_repo_root() -> Path:
    current = Path(__file__).resolve()
    for candidate in [current.parent, *current.parents]:
        if (candidate / ".git").exists():
            return candidate
    raise RuntimeError("Nem található a repository gyökere (.git).")


def resolve_file(repo_root: Path, expected: list[Path], pattern: str, description: str) -> Path:
    for path in expected:
        if path.exists():
            return path

    matches = sorted(p for p in repo_root.rglob(pattern) if p.is_file())
    if not matches:
        expected_text = "\n - ".join(str(p) for p in expected)
        raise FileNotFoundError(
            f"Hiányzó bemenet: {description}. Kerestem itt:\n - {expected_text}\n"
            f"és globbal: {pattern}."
        )
    if len(matches) > 1:
        match_text = "\n - ".join(str(p) for p in matches)
        raise RuntimeError(
            f"Több lehetséges fájl található ehhez: {description}. Pontosítsd az elérési utat.\n - {match_text}"
        )
    return matches[0]


def select_column(columns: list[str], include_terms: list[str], optional_terms: list[str] | None = None) -> str:
    optional_terms = optional_terms or []
    normalized = {col: normalize_text(col) for col in columns}

    for col, key in normalized.items():
        if all(term in key for term in include_terms) and all(term in key for term in optional_terms):
            return col

    for col, key in normalized.items():
        if all(term in key for term in include_terms):
            return col

    raise KeyError(f"Nem található oszlop: kötelező elemek={include_terms}, opcionális={optional_terms}")


def select_ksh_column(columns: list[str]) -> str:
    normalized = {col: normalize_text(col) for col in columns}
    for col, key in normalized.items():
        if (
            ("ksh" in key and ("kod" in key or "azonosito" in key or "torzsszam" in key))
            or ("telepules" in key and "azonosito" in key)
        ):
            return col
    raise KeyError("Nem található KSH/településazonosító oszlop.")


def detect_ksh_sheet_and_header(excel_path: Path) -> tuple[str, int]:
    """Find the sheet and header row by scanning for the row containing KSH code + population fields."""
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            for row_idx in range(1, min(sheet.max_row, 60) + 1):
                row_values = [normalize_text(cell.value) for cell in sheet[row_idx] if cell.value is not None]
                if not row_values:
                    continue

                has_ksh = any(
                    ("ksh" in value and ("kod" in value or "azonosito" in value or "torzsszam" in value))
                    or ("telepules" in value and "azonosito" in value)
                    for value in row_values
                )
                has_population = any("lako" in value and "nepesseg" in value for value in row_values)
                has_county = any("megye" in value or "varmegye" in value for value in row_values)
                if has_ksh and has_population and has_county:
                    return sheet_name, row_idx
    finally:
        workbook.close()

    raise RuntimeError(
        f"Nem sikerült fejlécet találni a KSH fájlban: {excel_path}. "
        "Ellenőrizd, hogy tartalmazza-e a KSH kód, lakónépesség és megye/vármegye mezőket."
    )


def load_ksh_population(excel_path: Path) -> pd.DataFrame:
    sheet_name, header_row = detect_ksh_sheet_and_header(excel_path)
    raw_df = pd.read_excel(excel_path, sheet_name=sheet_name, header=header_row - 1, dtype=str)

    raw_df = raw_df.dropna(how="all")
    columns = [str(col) for col in raw_df.columns]

    ksh_col = select_ksh_column(columns)
    pop_col = select_column(columns, ["lako", "nepesseg"])
    county_col = select_column(columns, ["megye"])
    name_col = select_column(columns, ["helyseg", "megnevez"])

    df = raw_df[[ksh_col, pop_col, county_col, name_col]].copy()
    df.columns = ["ksh_code_raw", "population_raw", "county", "settlement_name_ksh"]

    df["ksh_code"] = df["ksh_code_raw"].map(normalize_ksh_code)
    df["population"] = pd.to_numeric(df["population_raw"], errors="coerce")
    df["county"] = df["county"].astype("string").str.strip()
    df["settlement_name_norm"] = df["settlement_name_ksh"].map(normalize_name)

    df = df[df["ksh_code"].notna()].copy()
    df = df.drop_duplicates(subset=["ksh_code"], keep="first")

    return df[["ksh_code", "population", "county", "settlement_name_ksh", "settlement_name_norm"]]


def ensure_valid_geodata(gdf: gpd.GeoDataFrame, path: Path, expected_geom: str) -> None:
    if gdf.empty:
        raise ValueError(f"Üres állomány: {path}")
    if gdf.crs is None:
        raise ValueError(f"Hiányzó CRS a fájlban: {path}. Állíts be vetületet az inputban.")
    if gdf.geometry.is_empty.any():
        raise ValueError(f"Üres geometriák találhatók a fájlban: {path}")
    if gdf.geometry.isna().any():
        raise ValueError(f"Hiányzó geometriák találhatók a fájlban: {path}")

    geom_types = set(gdf.geometry.geom_type.unique())
    if expected_geom == "point" and not geom_types <= {"Point", "MultiPoint"}:
        raise ValueError(f"A boltok geometriája pont kell legyen. Talált típusok: {sorted(geom_types)}")


def safe_float(value: object) -> float | None:
    if value is None:
        return None
    if isinstance(value, (np.floating, float)):
        if math.isnan(float(value)):
            return None
        return float(value)
    if isinstance(value, (np.integer, int)):
        return float(value)
    return None


def make_json_serializable(obj: object) -> object:
    if isinstance(obj, dict):
        return {k: make_json_serializable(v) for k, v in obj.items()}
    if isinstance(obj, list):
        return [make_json_serializable(v) for v in obj]
    if isinstance(obj, (np.integer, int)):
        return int(obj)
    if isinstance(obj, (np.floating, float)):
        if math.isnan(float(obj)):
            return None
        return float(obj)
    if obj is pd.NA:
        return None
    return obj


def main() -> None:
    repo_root = find_repo_root()

    shops_path = resolve_file(
        repo_root,
        [repo_root / "data" / "shops.geojson", repo_root / "3. házi" / "elelmiszersivatagok" / "data" / "shops.geojson"],
        "shops.geojson",
        "boltok GeoJSON",
    )
    settlements_path = resolve_file(
        repo_root,
        [repo_root / "data" / "settlements.geojson", repo_root / "3. házi" / "elelmiszersivatagok" / "data" / "settlements.geojson"],
        "settlements.geojson",
        "települések GeoJSON",
    )
    ksh_path = resolve_file(
        repo_root,
        [repo_root / "3. házi" / "hnt_letoltes_2026.xlsx"],
        "hnt_letoltes_*.xlsx",
        "KSH Helységnévtár Excel",
    )

    output_data_dir = repo_root / "data"
    output_data_dir.mkdir(parents=True, exist_ok=True)
    results_path = output_data_dir / "results.geojson"
    summary_path = output_data_dir / "summary.json"

    shops = gpd.read_file(shops_path)
    settlements = gpd.read_file(settlements_path)
    ensure_valid_geodata(shops, shops_path, expected_geom="point")
    ensure_valid_geodata(settlements, settlements_path, expected_geom="polygon")

    if "ksh_ref" not in settlements.columns:
        raise KeyError(f"A település állományban nincs 'ksh_ref' oszlop: {settlements_path}")
    if "name" not in settlements.columns:
        raise KeyError(f"A település állományban nincs 'name' oszlop: {settlements_path}")

    if settlements.crs.to_string() != shops.crs.to_string():
        shops = shops.to_crs(settlements.crs)

    ksh = load_ksh_population(ksh_path)

    settlements = settlements.copy()
    settlements["ksh_code"] = settlements["ksh_ref"].map(normalize_ksh_code)
    settlements["settlement_name_norm"] = settlements["name"].map(normalize_name)

    merged = settlements.merge(ksh[["ksh_code", "population", "county"]], how="left", on="ksh_code")
    unmatched_by_code = merged["population"].isna().sum()
    print(f"Nem egyező települések KSH azonosító alapján: {int(unmatched_by_code)}")

    # Fallback by settlement name for still-unmatched rows (only unique names are used to avoid ambiguous matches).
    unmatched_mask = merged["population"].isna()
    ksh_name_unique = ksh.dropna(subset=["settlement_name_norm"]).copy()
    dup_name_mask = ksh_name_unique["settlement_name_norm"].duplicated(keep=False)
    ksh_name_unique = ksh_name_unique.loc[~dup_name_mask, ["settlement_name_norm", "population", "county"]]

    fallback_map = ksh_name_unique.set_index("settlement_name_norm")
    for idx in merged.index[unmatched_mask]:
        key = merged.at[idx, "settlement_name_norm"]
        if key in fallback_map.index:
            merged.at[idx, "population"] = fallback_map.at[key, "population"]
            merged.at[idx, "county"] = fallback_map.at[key, "county"]

    remaining_unmatched = merged["population"].isna().sum()
    print(f"Nem egyező települések név szerinti pótlás után: {int(remaining_unmatched)}")

    shop_join = gpd.sjoin(shops[["geometry"]], merged[["geometry"]], how="left", predicate="within")
    shop_counts = shop_join.groupby("index_right").size()
    merged["shop_count"] = merged.index.to_series().map(shop_counts).fillna(0).astype(int)

    merged["residents_per_shop"] = np.where(
        merged["shop_count"] > 0,
        merged["population"] / merged["shop_count"],
        np.nan,
    )

    settlements_eov = merged.to_crs(epsg=EOV_EPSG)
    shops_eov = shops.to_crs(epsg=EOV_EPSG)

    if shops_eov.empty:
        raise ValueError("A boltok állománya üres, ezért nem számolható legközelebbi bolt távolság.")

    shop_union = shops_eov.unary_union
    centroids = settlements_eov.geometry.centroid
    no_shop_mask = merged["shop_count"] == 0

    distances_km = pd.Series(np.nan, index=merged.index, dtype="float64")
    distances_km.loc[no_shop_mask] = centroids.loc[no_shop_mask].distance(shop_union) / 1000.0
    merged["nearest_shop_km"] = distances_km

    merged["category"] = np.select(
        [
            merged["shop_count"] > 0,
            no_shop_mask & (merged["nearest_shop_km"] <= 5),
            no_shop_mask & (merged["nearest_shop_km"] <= 10),
        ],
        [
            "van bolt",
            "nincs bolt, 5 km-en belül van",
            "nincs bolt, 5–10 km",
        ],
        default="nincs bolt, 10 km felett",
    )

    no_shop = merged[merged["shop_count"] == 0].copy()
    no_shop["population"] = pd.to_numeric(no_shop["population"], errors="coerce")

    category_summary = []
    for category, group in merged.groupby("category", dropna=False):
        category_summary.append(
            {
                "category": category,
                "settlement_count": int(len(group)),
                "resident_count": float(group["population"].sum(min_count=1)) if group["population"].notna().any() else None,
            }
        )

    county_breakdown = []
    for county, group in no_shop.groupby("county", dropna=False):
        county_category_counts = {
            cat: int((group["category"] == cat).sum())
            for cat in [
                "nincs bolt, 5 km-en belül van",
                "nincs bolt, 5–10 km",
                "nincs bolt, 10 km felett",
            ]
        }
        county_breakdown.append(
            {
                "county": None if pd.isna(county) else str(county),
                "settlement_count": int(len(group)),
                "resident_count": float(group["population"].sum(min_count=1)) if group["population"].notna().any() else None,
                "categories": county_category_counts,
            }
        )

    summary = {
        "inputs": {
            "shops": str(shops_path),
            "settlements": str(settlements_path),
            "ksh": str(ksh_path),
        },
        "totals": {
            "settlement_count": int(len(merged)),
            "matched_population_count": int(merged["population"].notna().sum()),
            "unmatched_population_count": int(merged["population"].isna().sum()),
            "resident_count_total": float(merged["population"].sum(min_count=1)) if merged["population"].notna().any() else None,
            "settlements_without_shop": int(len(no_shop)),
            "residents_without_shop": float(no_shop["population"].sum(min_count=1)) if no_shop["population"].notna().any() else None,
        },
        "categories": category_summary,
        "county_breakdown_without_shop": sorted(county_breakdown, key=lambda x: ("" if x["county"] is None else x["county"])),
    }

    merged = merged.to_crs(OUTPUT_CRS)
    merged["population"] = merged["population"].map(safe_float)
    merged["residents_per_shop"] = merged["residents_per_shop"].map(safe_float)
    merged["nearest_shop_km"] = merged["nearest_shop_km"].map(safe_float)

    merged.to_file(results_path, driver="GeoJSON")
    with summary_path.open("w", encoding="utf-8") as f:
        json.dump(make_json_serializable(summary), f, ensure_ascii=False, indent=2)

    print(f"Mentve: {results_path}")
    print(f"Mentve: {summary_path}")


if __name__ == "__main__":
    main()
